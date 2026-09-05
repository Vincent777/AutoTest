from openpyxl import load_workbook
import re
import os
import sys


PRESET_ENV_FILTER = [
    "CUDA_VISIBLE_DEVICES",
    "SIG_LOG_LEVEL",
    "SIGINFER_C",
    "LD_LIBRARY_PATH",
]

# excel/docker 行里可能出现的镜像前缀
IMAGE_PREFIX_RE = (
    r"^.*(?:"
    r"(?:[\w.-]+/)*siginfer-(?:aarch64|x86_64)-iluvatar|"
    r"docker\.xcoresigma\.com/docker/siginfer[^:\s]+|"
    r"xcoresigma-registry\.cn-beijing\.cr\.aliyuncs\.com/docker/siginfer[^:\s]+|"
    r"docker\.infinitensor\.com/docker/siginfer[^:\s]+"
    r"):\S+\s*"
)


def normalize_iluvatar_launch_args(result: str) -> str:
    """将历史参数归一为 Iluvatar 现场可用参数（对齐手动 docker run）。"""
    result = re.sub(r"--platform-type\s+\S+", "--platform-type iluvatar", result)
    # -tp N / --tp N -> --tensor-parallel-size N
    result = re.sub(r"(?:^|\s)-(-)?tp(?:\s+|=)(\d+)", r" --tensor-parallel-size \2", result)
    # --weight-dtype=FP16 -> --weight-dtype FP16
    result = re.sub(r"--weight-dtype\s*=\s*(\S+)", r"--weight-dtype \1", result)
    # tokenizer 路径补尾斜杠
    tok_m = re.search(r"--tokenizer\s+(\S+)", result)
    if tok_m and not tok_m.group(1).endswith("/"):
        result = re.sub(
            r"--tokenizer\s+" + re.escape(tok_m.group(1)),
            f"--tokenizer {tok_m.group(1)}/",
            result,
            count=1,
        )
    result = re.sub(r"\s+", " ", result).strip()
    if result and not result.startswith(" "):
        result = " " + result
    return result


def main():
    if len(sys.argv) != 3:
        print("Usage: python script_generator_for_SigInfer.py <test_type> <version>")
        sys.exit(1)

    test_type = sys.argv[1]
    version = sys.argv[2]

    curr_dir = os.getcwd()

    file_path = f'{curr_dir}/{version}/SigInfer_model_list.xlsx'
    workbook = load_workbook(file_path)

    if 'Iluvatar' in workbook.sheetnames:
        sheet = workbook['Iluvatar']
    else:
        sheet = workbook.active

    row_count = sheet.max_row
    print(f"总行数: {row_count}")

    if test_type == "Smoke":
        target_file = "SigInfer_job_executor_for_SmokeTest.sh"
    elif test_type == "Performance":
        target_file = "SigInfer_job_executor_for_PerformanceTest.sh"
    elif test_type == "Stability":
        target_file = "SigInfer_job_executor_for_StabilityTest.sh"
    elif test_type == "Accuracy":
        target_file = "SigInfer_job_executor_for_AccuracyTest.sh"
    else:
        print(f"Unsupported test_type: {test_type}")
        sys.exit(1)

    src_code = ""
    start = True
    env_vars = []
    for row in sheet.iter_rows(min_row=2, max_row=row_count, values_only=True):
        name = row[0]
        args = row[2]
        if not name or not args:
            continue
        args = str(args).split('\n')[0]

        env_vars += [
            v for v in re.findall(r'-e\s+"?([^"\s]+=[^"\s]+)"?', args)
            if v.split('=')[0] not in PRESET_ENV_FILTER
        ]

        result = re.sub(IMAGE_PREFIX_RE, "", args)
        result = re.sub(
            r"^docker\s+run\b.*?(?=\s--model\b|\s-tp\b|\s--tokenizer\b|\s--tensor-parallel-size\b)",
            "",
            result,
        )
        result = re.sub(r"--swap-space\s+\d+", "$SWAP_SPACE_OPTION", result)
        result = re.sub(r"--prometheus-port\s+\d+", "--prometheus-port $PROMETHEUS_PORT", result)
        result = re.sub(r"--port\s+\d+", "--port $PORT", result)
        result = re.sub(r"--master-addr\s+\S+", "--master-addr $MASTER_IP", result)
        result = re.sub(r"--node-rank\s+\d+", "--node-rank $NODE_RANK", result)
        if "--schedule-policy" in result:
            result = re.sub(r"--schedule-policy\s+\S+", "--schedule-policy $SCHEDULE_POLICY", result)
        else:
            result += " --schedule-policy $SCHEDULE_POLICY"
        result = re.sub(r"--master-port\s+\d+", "--master-port $MASTER_PORT", result)
        if "--use-prefix-cache" in result:
            result = re.sub(r"--use-prefix-cache", "$USE_PREFIX_CACHE", result)
        else:
            result += " $USE_PREFIX_CACHE"
        if "--prometheus-port" not in result:
            result += " --prometheus-port $PROMETHEUS_PORT"
        result = normalize_iluvatar_launch_args(result)

        if start:
            src_code += f"if [ $MODEL == \"{name}\" ]; then\n"
            start = False
        else:
            src_code += f"elif [ $MODEL == \"{name}\" ]; then\n"
        src_code += "    echo \"SigInfer"
        src_code += result
        src_code += "\"\n"

        # 镜像默认 ENTRYPOINT 直接接收启动参数（对齐手动 docker run）
        src_code += "    EXEC_COMMAND+=\""
        src_code += result
        src_code += " > $LOG_NAME 2>&1 &\"\n"

    src_code += "fi\n"
    env_vars = list(dict.fromkeys(env_vars))

    template_file = "job_executor_template_for_SigInfer.sh"

    try:
        with open(f"{curr_dir}/{template_file}", 'r', encoding='utf-8') as file:
            lines = file.readlines()

        line_num = 0
        for line in lines:
            if "<<<generated source code>>>" in line:
                lines[line_num] = src_code
                with open(f"{curr_dir}/{target_file}", 'w') as file:
                    file.write(''.join(lines))
                break
            elif "<<<TEST_TYPE>>>" in line:
                mapping = {
                    "Smoke": "SmokeTest",
                    "Performance": "PerformanceTest",
                    "Accuracy": "AccuracyTest",
                    "Stability": "StabilityTest",
                }
                lines[line_num] = line.replace("<<<TEST_TYPE>>>", mapping[test_type])
            elif "<<<ENV_VARS>>>" in line:
                if len(env_vars) > 0:
                    env_var_lines = ""
                    for var_def in env_vars:
                        env_var_lines += f"     -e {var_def} \\\n"
                    lines[line_num] = env_var_lines
                else:
                    lines[line_num] = ""
            line_num += 1
    except FileNotFoundError:
        print(f"Error: template '{curr_dir}/{template_file}' not found.")
        sys.exit(1)
    except Exception as e:
        print(f"Error reading file: {str(e)}")
        sys.exit(1)

    os.system(f"chmod 777 {target_file}")
    print(f"Generated: {target_file}")


if __name__ == "__main__":
    main()
