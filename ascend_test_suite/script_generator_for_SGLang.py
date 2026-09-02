from openpyxl import load_workbook
import re
import os
import sys


def extract_tp(args: str) -> str:
    for pattern in (r"--tp-size\s+(\d+)", r"-tp\s+(\d+)", r"--tensor-parallel-size\s+(\d+)"):
        match = re.search(pattern, args)
        if match:
            return match.group(1)
    return "1"


def extract_model_path(name: str, args: str) -> str:
    patterns = (
        r"--model-path\s+(\S+)",
        r"--tokenizer\s+(\S+)",
        r"--model\s+(\S+)",
    )
    for pattern in patterns:
        match = re.search(pattern, args)
        if match:
            path = match.group(1).rstrip("/")
            if path.startswith("/"):
                return path
            return f"/home/weight/{path}"

    if name.startswith("Qwen3-"):
        if name == "Qwen3-235B-A22B":
            return "/home/weight/Qwen3/Qwen3-235B-A22B"
        if name == "Qwen3-32B-FP8":
            return "/home/weight/Qwen3/Qwen3-32B-FP8"
        return f"/home/weight/Qwen3/{name}"
    return f"/home/weight/{name}"


def extract_optional_flag(args: str, names: tuple[str, ...]) -> str:
    for flag in names:
        match = re.search(rf"{re.escape(flag)}\s+(\S+)", args)
        if match:
            return f"{flag} {match.group(1)}"
    return ""


def strip_docker_prefix(args: str) -> str:
    patterns = (
        r"^.*docker\.xcoresigma\.com/docker/\S+",
        r"^.*quay\.io/ascend/sglang(?:-ascend)?:\S+",
    )
    result = args
    for pattern in patterns:
        result = re.sub(pattern, "", result)
    return result.strip()


def normalize_sglang_args(name: str, args: str) -> str:
    args = (args or "").split("\n")[0]
    result = strip_docker_prefix(args)

    # DSV4 Pagoda / sglang-universal-plugin: keep run_dsv4_*.sh, do not rewrite to launch_server.
    # 禁止在容器内用 setsid：setsid 会 fork 后父进程退出，Docker 把 PID1 结束当成容器退出 (Exited 0)。
    if re.search(r"run_dsv4", args, re.I):
        m = re.search(r"((?:setsid\s+)?bash\s+\S*run_dsv4\S*)", result, re.I)
        if m:
            result = m.group(1).strip()
        result = re.sub(r"^setsid\s+", "", result, flags=re.I)
        if not re.match(r"bash\s+", result, re.I):
            result = f"bash {result}"
        # CI mounts host /home/s_limingge; prefer absolute script path
        result = re.sub(
            r"(bash\s+)(?!/)(\S*run_dsv4\S*)",
            r"\1/home/s_limingge/sglang-universal-plugin/\2",
            result,
            flags=re.I,
        )
        # drop excel-only flags / leftover args; port is positional $PORT
        result = re.sub(r"--tp-size\s+\d+", "", result)
        result = re.sub(r"--port\s+\S+", "", result)
        result = re.sub(r"\s+\d+\s*$", "", result)
        # strip anything after the script path
        result = re.sub(r"(bash\s+\S*run_dsv4\S*).*", r"\1", result, flags=re.I)
        script = result.strip()
        # entry_points 需 pip install -e；同时强制 PYTHONPATH，避免 editable 元数据在却 import 失败
        # 注意：result 会嵌入 EXEC_COMMAND+="..."，故 bash -lc 外层用 \"，内层 python -c 用单引号
        plugin_root = "/home/s_limingge/sglang-universal-plugin"
        result = (
            "bash -lc "
            "\\"
            '"'
            f"set -e; "
            f"cd {plugin_root}; "
            f"export PYTHONPATH={plugin_root}/src; "
            f"pip install -e . --no-deps; "
            f"python3 -c 'import sglang_universal_plugin'; "
            f"{script} $PORT"
            "\\"
            '"'
        )
        return re.sub(r"\s+", " ", result).strip()

    # 1) 已是标准入口则沿用；2) 仅有 *launch_server 则规范化；3) 都没有则按模型拼默认命令
    if "sglang.launch_server" not in result:
        rewritten = re.sub(r"^.*launch_server", "python3 -m sglang.launch_server", result)
        if "sglang.launch_server" in rewritten:
            result = rewritten
        else:
            model_path = extract_model_path(name, args)
            tp = extract_tp(args)
            context = extract_optional_flag(args, ("--context-length", "--max-model-len"))
            mem = extract_optional_flag(args, ("--mem-fraction-static", "--gpu-memory-utilization"))
            radix = ""
            if "--disable-radix-cache" not in args and "--enable-radix-cache" not in args:
                if re.search(r"--no-enable-prefix-caching|prefix.cach", args, re.I):
                    radix = " --disable-radix-cache"
            parts = [
                "python3 -m sglang.launch_server",
                f"--model-path {model_path}",
                f"--served-model-name {name}",
                "--port 4321",
                f"--tp-size {tp}",
                "--host 0.0.0.0",
            ]
            if context:
                if context.startswith("--max-model-len"):
                    context = context.replace("--max-model-len", "--context-length")
                parts.append(context)
            if mem:
                if mem.startswith("--gpu-memory-utilization"):
                    mem = mem.replace("--gpu-memory-utilization", "--mem-fraction-static")
                parts.append(mem)
            if radix:
                parts.append(radix.strip())
            elif "--disable-radix-cache" in args:
                parts.append("--disable-radix-cache")
            result = " ".join(parts)

    result = re.sub(r"--port\s+\d+", "--port $PORT", result)
    if not re.search(r"--port\s+\$PORT\b|--port\s+\d+", result):
        result = f"{result} --port $PORT"
    result = re.sub(r"--served-model-name\s+\S+", f"--served-model-name {name}", result)
    # Excel 合部命令里若已手写 PD 参数，去掉以免和 $PD_EXTRA_ARGS 重复
    result = re.sub(r"--disaggregation-mode\s+\S+", "", result)
    result = re.sub(r"--disaggregation-transfer-backend\s+\S+", "", result)
    result = re.sub(r"--disaggregation-ib-device\s+\S+", "", result)
    result = re.sub(r"--disaggregation-bootstrap-port\s+\S+", "", result)
    # Decode 侧会注入 --max-running-requests / --cuda-graph-max-bs，避免 argparse 重复
    result = re.sub(r"--max-running-requests\s+\S+", "", result)
    result = re.sub(r"--cuda-graph-max-bs(?:-decode)?\s+\S+", "", result)
    result = re.sub(r"--disable-overlap-schedule\b", "", result)
    return re.sub(r"\s+", " ", result).strip()


def main():
    if len(sys.argv) != 3:
        print("Usage: python script_generator_for_SGLang.py <test_type> <version>")
        sys.exit(1)

    test_type = sys.argv[1]
    version = sys.argv[2]
    curr_dir = os.getcwd()
    file_path = f"{curr_dir}/{version}/SGLang_model_list.xlsx"
    workbook = load_workbook(file_path)
    
    sheet = workbook["Ascend"]
    row_count = sheet.max_row
    print(f"总行数: {row_count}")

    if test_type == "Smoke":
        target_file = "SGLang_job_executor_for_SmokeTest.sh"
    elif test_type == "Performance":
        target_file = "SGLang_job_executor_for_PerformanceTest.sh"
    elif test_type == "Stability":
        target_file = "SGLang_job_executor_for_StabilityTest.sh"
    elif test_type == "Accuracy":
        target_file = "SGLang_job_executor_for_AccuracyTest.sh"
    else:
        print(f"Unsupported test_type: {test_type}")
        sys.exit(1)

    src_code = ""
    start = True
    for row in sheet.iter_rows(min_row=2, max_row=row_count, values_only=True):
        name = row[0]
        if not name:
            continue
        args = row[2]
        result = normalize_sglang_args(name, args)

        if start:
            src_code += f'if [ $MODEL == "{name}" ]; then\n'
            start = False
        else:
            src_code += f'elif [ $MODEL == "{name}" ]; then\n'
        # run_dsv4 脚本只吃位置参数端口，不要追加 launch_server 专用 flag
        if re.search(r"run_dsv4", result, re.I):
            src_code += f'    echo "{result}"\n'
            src_code += (
                f'    EXEC_COMMAND+=" {result} > $LOG_NAME 2>&1 &"\n'
            )
        else:
            src_code += f'    echo "{result} $PD_EXTRA_ARGS"\n'
            src_code += (
                f'    EXEC_COMMAND+=" $PD_DOCKER_CMD_PREFIX {result} '
                f'$SGLANG_PREFIX_CACHE $PD_EXTRA_ARGS $PD_DOCKER_CMD_SUFFIX > $LOG_NAME 2>&1 &"\n'
            )

    src_code += "fi\n"

    template_file = "job_executor_template_for_SGLang.sh"
    try:
        with open(f"{curr_dir}/{template_file}", "r", encoding="utf-8") as file:
            lines = file.readlines()

        line_num = 0
        for line in lines:
            if "<<<generated source code>>>" in line:
                lines[line_num] = src_code
                with open(f"{curr_dir}/{target_file}", "w", encoding="utf-8") as file:
                    file.write("".join(lines))
                break
            if "<<<TEST_TYPE>>>" in line:
                if test_type == "Smoke":
                    lines[line_num] = line.replace("<<<TEST_TYPE>>>", "SmokeTest")
                elif test_type == "Performance":
                    lines[line_num] = line.replace("<<<TEST_TYPE>>>", "PerformanceTest")
                elif test_type == "Accuracy":
                    lines[line_num] = line.replace("<<<TEST_TYPE>>>", "AccuracyTest")
                elif test_type == "Stability":
                    lines[line_num] = line.replace("<<<TEST_TYPE>>>", "StabilityTest")
            line_num += 1
    except FileNotFoundError:
        print(f"Error: template file '{curr_dir}/{template_file}' not found.")
        sys.exit(1)
    except Exception as e:
        print(f"Error reading file: {str(e)}")
        sys.exit(1)

    os.system(f"chmod 777 {target_file}")


if __name__ == "__main__":
    main()
