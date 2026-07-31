#!/usr/bin/env python3
"""Convert ascend_test_suite Performance Excel reports to unified JSON and save locally.

Example:
  python3 collector/excel_to_json.py \\
    ../ascend_test_suite/report_20260727/11111/DeepSeek-R1-Distill-Qwen-32B_*.xlsx \\
    ../ascend_test_suite/report_20260727/00000/DeepSeek-R1-Distill-Qwen-32B_*.xlsx

  # optional ingest into SQLite
  python3 collector/excel_to_json.py ... --ingest
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT_DIR = ROOT / "results"


def detect_engine(launch_cmd: str) -> str:
    """Fallback only: infer engine from launch command text."""
    text = (launch_cmd or "").lower()
    if "sglang" in text:
        return "sglang"
    if "vllm" in text:
        return "vllm"
    if "mindie" in text:
        return "mindie"
    if "siginfer" in text:
        return "siginfer"
    return "unknown"


def normalize_engine_name(raw: Any) -> str | None:
    """Normalize Excel 'Inference Engine' cell to canonical lowercase id."""
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    key = re.sub(r"[\s_\-]+", "", text.lower())
    mapping = {
        "siginfer": "siginfer",
        "vllm": "vllm",
        "sglang": "sglang",
        "mindie": "mindie",
    }
    if key in mapping:
        return mapping[key]
    # Keep unknown labels as lowercase slug for visibility
    return key


def read_inference_engine_from_excel(ws) -> str | None:
    """Read engine from the Inference Engine column (usually col 1, data from row 6)."""
    engine_col = 1
    # Prefer header on row 4 (merged with row 5 in template)
    for col in range(1, (ws.max_column or 8) + 1):
        for row in (4, 5):
            val = ws.cell(row, col).value
            if val and "inference engine" in str(val).strip().lower():
                engine_col = col
                break

    for row in range(6, ws.max_row + 1):
        eng = normalize_engine_name(ws.cell(row, engine_col).value)
        if eng:
            return eng
    return None


def detect_engine_version(launch_cmd: str, version_file: Path | None) -> str:
    if version_file and version_file.exists():
        ver = version_file.read_text(encoding="utf-8").strip()
        if ver:
            return ver
    patterns = (
        r"quay\.io/ascend/sglang(?:-ascend)?:([^\s]+)",
        r"quay\.io/ascend/vllm-ascend:([^\s]+)",
        r"sglang:([^\s]+)",
        r"vllm-ascend:([^\s]+)",
        r"vllm/vllm-openai:([^\s]+)",
    )
    for pat in patterns:
        m = re.search(pat, launch_cmd or "")
        if m:
            return m.group(1)
    return "unknown"


def clean_model_name(raw: str) -> str:
    name = (raw or "").strip()
    for suffix in (
        "_DynamicSplitFuseV2_Use-prefix-cache_Swap-space",
        "_DynamicSplitFuseV2_Use-prefix-cache",
        "_DynamicSplitFuseV2_Swap-space",
        "_DynamicSplitFuseV2",
        "_Use-prefix-cache_Swap-space",
        "_Use-prefix-cache",
        "_Swap-space",
    ):
        if name.endswith(suffix):
            return name[: -len(suffix)]
    return name


def parse_batch_cell(value: Any) -> tuple[int | None, int | None]:
    """Parse '1 (num_prompt=4)' -> (concurrency=1, num_prompts=4)."""
    if value is None:
        return None, None
    text = str(value)
    m = re.match(r"\s*(\d+)\s*\(num_prompt=(\d+)\)", text)
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.match(r"\s*(\d+)\s*$", text)
    if m:
        return int(m.group(1)), None
    return None, None


def parse_context_cell(value: Any) -> tuple[int | None, int | None, str]:
    """Parse Excel 「上下文长度」 -> (input_len, output_len, workload_id).

    workload_id keeps the Excel cell text as-is (e.g. ``128+128``, ``1024+4096``,
    ``SharedGPT``). Do not rewrite digits to ``1k`` or add ``isl``/``osl`` prefixes.
    """
    if value is None:
        return None, None, "unknown"
    text = str(value).strip()
    m = re.match(r"(\d+)\+(\d+)$", text)
    if m:
        return int(m.group(1)), int(m.group(2)), text
    return None, None, text


def report_date_from_path(path: Path) -> str | None:
    for part in path.parts:
        m = re.match(r"report_(\d{8})", part)
        if m:
            return m.group(1)
    return None


def parse_perf_log_request_throughput(text: str) -> dict[tuple[str, int], dict[str, float]]:
    """Parse Serving Benchmark Result blocks keyed by (context, concurrency)."""
    results: dict[tuple[str, int], dict[str, float]] = {}
    in_out_length_key: str | None = None
    concurrency_key: int | None = None
    matches = re.finditer(
        r"(Random Testing [^\n]+|Testing concurrency=[^\n]+|[\=]+ Serving Benchmark Result [\=]+.*?[\=]+)",
        text,
        re.DOTALL,
    )
    for match in matches:
        section = match.group(0)
        io = re.search(r"Random Testing input=(\d+), output=(\d+)", section)
        if io:
            in_out_length_key = f"{io.group(1)}+{io.group(2)}"
        cc = re.search(r"Testing concurrency=(\d+), prompts=(\d+)", section)
        if cc:
            concurrency_key = int(cc.group(1))
            if in_out_length_key is None:
                in_out_length_key = "SharedGPT"
            results[(in_out_length_key, concurrency_key)] = {}
        bm = re.search(r"[\=]+ Serving Benchmark Result [\=]+.*?[\=]+", section, re.DOTALL)
        if bm and in_out_length_key is not None and concurrency_key is not None:
            key = (in_out_length_key, concurrency_key)
            results.setdefault(key, {})
            for metric in re.finditer(
                r"(?P<key>[A-Za-z0-9\s\(\)/]+):\s+(?P<value>[0-9.]+)",
                bm.group(0),
            ):
                results[key][metric.group("key").strip()] = float(metric.group("value"))
    return results


def find_header_columns(ws) -> dict[str, int]:
    """Map row-5 header names to 1-based column indexes."""
    mapping: dict[str, int] = {}
    for col in range(1, (ws.max_column or 20) + 1):
        val = ws.cell(5, col).value
        if not val:
            continue
        text = str(val).strip().lower()
        mapping[text] = col
    return mapping


def col_by_alias(headers: dict[str, int], *aliases: str) -> int | None:
    for alias in aliases:
        if alias.lower() in headers:
            return headers[alias.lower()]
    return None


def excel_to_records(
    excel_path: Path,
    *,
    hardware: str = "Ascend910B",
    engine: str | None = None,
    engine_version: str | None = None,
    perf_log: Path | None = None,
) -> list[dict[str, Any]]:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active
    headers = find_header_columns(ws)

    model_raw = ws.cell(1, 2).value or excel_path.stem
    model = clean_model_name(str(model_raw))
    launch_cmd = str(ws.cell(2, 2).value or "")
    test_cmd = str(ws.cell(3, 2).value or "")

    version_file = excel_path.parent / "version.txt"
    # Primary: Excel "Inference Engine" column; optional CLI override; launch_cmd as last resort
    eng = engine or read_inference_engine_from_excel(ws) or detect_engine(launch_cmd)
    ver = engine_version or detect_engine_version(launch_cmd, version_file)

    report_date = report_date_from_path(excel_path)
    started_at = None
    if report_date:
        started_at = datetime.strptime(report_date, "%Y%m%d").replace(tzinfo=timezone.utc).isoformat()

    # Optional enrichment from raw performance log (when Excel lacks Request throughput)
    log_metrics: dict[tuple[str, int], dict[str, float]] = {}
    if perf_log and perf_log.exists():
        log_metrics = parse_perf_log_request_throughput(
            perf_log.read_text(encoding="utf-8", errors="replace")
        )

    col_success = col_by_alias(headers, "Successful requests") or 4
    col_req = col_by_alias(headers, "Request throughput")
    col_out = col_by_alias(headers, "Output token throughput") or 5
    col_total = col_by_alias(headers, "Total Token throughput", "Total token throughput")
    col_ttft_mean = col_by_alias(headers, "Mean TTFT")
    col_ttft_p50 = col_by_alias(headers, "Median TTFT")
    col_ttft_p99 = col_by_alias(headers, "P99 TTFT")
    col_tpot_mean = col_by_alias(headers, "Mean TPOT")
    col_tpot_p50 = col_by_alias(headers, "Median TPOT")
    col_tpot_p99 = col_by_alias(headers, "P99 TPOT")
    col_itl_mean = col_by_alias(headers, "Mean ITL")
    col_itl_p50 = col_by_alias(headers, "Median ITL")
    col_itl_p99 = col_by_alias(headers, "P99 ITL")

    def cell_float(r: int, col: int | None) -> float | None:
        if not col:
            return None
        val = ws.cell(r, col).value
        if val is None:
            return None
        try:
            return float(val)
        except (TypeError, ValueError):
            return None

    records: list[dict[str, Any]] = []
    current_context = None

    for row in range(6, ws.max_row + 1):
        ctx_cell = ws.cell(row, 2).value
        if ctx_cell is not None:
            current_context = ctx_cell

        batch_cell = ws.cell(row, 3).value
        concurrency, num_prompts = parse_batch_cell(batch_cell)
        if concurrency is None:
            continue

        successful = cell_float(row, col_success)
        out_tps = cell_float(row, col_out)
        req_tps = cell_float(row, col_req)
        if successful is None and out_tps is None and req_tps is None:
            continue

        input_len, output_len, workload = parse_context_cell(current_context)
        mean_ttft = cell_float(row, col_ttft_mean)
        median_ttft = cell_float(row, col_ttft_p50)
        p99_ttft = cell_float(row, col_ttft_p99)
        mean_tpot = cell_float(row, col_tpot_mean)
        median_tpot = cell_float(row, col_tpot_p50)
        p99_tpot = cell_float(row, col_tpot_p99)
        mean_itl = cell_float(row, col_itl_mean)
        median_itl = cell_float(row, col_itl_p50)
        p99_itl = cell_float(row, col_itl_p99)
        total_tps = cell_float(row, col_total)

        # Fallback: pull Request throughput from raw log
        if req_tps is None and current_context is not None and log_metrics:
            key = (str(current_context), concurrency)
            req_tps = (log_metrics.get(key) or {}).get("Request throughput (req/s)")

        success_rate = None
        if successful is not None and num_prompts:
            try:
                success_rate = float(successful) / float(num_prompts)
            except (TypeError, ValueError, ZeroDivisionError):
                success_rate = None

        run_id = f"{eng}_{ver}_{model}_{workload}_c{concurrency}_{report_date or 'na'}"
        records.append(
            {
                "run_id": run_id,
                "engine": eng,
                "engine_version": ver,
                "model": model,
                "hardware": hardware,
                "workload": workload,
                "concurrency": concurrency,
                "input_len": input_len,
                "output_len": output_len,
                "num_prompts": num_prompts,
                "metrics": {
                    "successful_requests": successful,
                    "request_throughput": req_tps,
                    "output_token_throughput": out_tps,
                    "total_token_throughput": total_tps,
                    "ttft_mean_ms": mean_ttft,
                    "ttft_p50_ms": median_ttft,
                    "ttft_p99_ms": p99_ttft,
                    "tpot_mean_ms": mean_tpot,
                    "tpot_p50_ms": median_tpot,
                    "tpot_p99_ms": p99_tpot,
                    "itl_mean_ms": mean_itl,
                    "itl_p50_ms": median_itl,
                    "itl_p99_ms": p99_itl,
                    "success_rate": success_rate,
                },
                "started_at": started_at,
                "finished_at": None,
                "artifacts_path": str(excel_path),
                "source": {
                    "excel": str(excel_path),
                    "launch_cmd": launch_cmd,
                    "test_cmd": test_cmd,
                    "excel_model_label": str(model_raw),
                    "perf_log": str(perf_log) if perf_log else None,
                },
            }
        )

    return records


def write_json(records: list[dict[str, Any]], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(records, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("excel_files", nargs="+", help="Performance Excel report path(s)")
    parser.add_argument("--out-dir", default=str(DEFAULT_OUT_DIR), help="Directory for JSON output")
    parser.add_argument("--merged", default="", help="Optional single merged JSON path")
    parser.add_argument("--hardware", default="Ascend910B")
    parser.add_argument("--engine", default="", help="Force engine name (vllm/sglang)")
    parser.add_argument("--engine-version", default="", help="Force engine version tag")
    parser.add_argument("--ingest", action="store_true", help="Also ingest into storage/perf.db")
    parser.add_argument(
        "--db",
        default=str(ROOT / "storage" / "perf.db"),
        help="SQLite db path used with --ingest",
    )
    args = parser.parse_args()

    out_dir = Path(args.out_dir)
    db_path = Path(args.db)
    all_records: list[dict[str, Any]] = []

    for excel in args.excel_files:
        path = Path(excel).resolve()
        if not path.exists():
            raise FileNotFoundError(path)
        records = excel_to_records(
            path,
            hardware=args.hardware,
            engine=args.engine or None,
            engine_version=args.engine_version or None,
        )
        eng = records[0]["engine"] if records else "unknown"
        ver = records[0]["engine_version"] if records else "unknown"
        model = records[0]["model"] if records else path.stem
        out_name = f"{eng}_{ver}_{model}.json".replace("/", "_")
        out_path = out_dir / out_name
        write_json(records, out_path)
        print(f"wrote {len(records)} record(s) -> {out_path}")
        all_records.extend(records)

        if args.ingest and records:
            from ingest import connect, ingest_file  # type: ignore

            conn = connect(db_path)
            n = ingest_file(conn, out_path)
            print(f"ingested {n} into {db_path}")
            conn.close()

    if args.merged:
        merged_path = Path(args.merged)
        write_json(all_records, merged_path)
        print(f"merged {len(all_records)} record(s) -> {merged_path}")

    print(f"total_records={len(all_records)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
